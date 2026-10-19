"""Build a test upload for the FPSO map: 15 projects from dataset2 + 5 made-up ones, no TAM/SAM/SOM columns.

Usage: python3 -I make_test_file.py <dataset2.xlsx> <output.xlsx>
"""
import html, re, sys, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

src, out = sys.argv[1], sys.argv[2]


def read_rows(path):
    """Read the first sheet straight from the XML (openpyxl returns empty cells for this file)."""
    z = zipfile.ZipFile(path)
    ss = [html.unescape(re.sub(r"<[^>]+>", "", s)) for s in
          re.findall(r"<si>(.*?)</si>", z.read("xl/sharedStrings.xml").decode(), re.S)]
    letters = lambda ref: re.match(r"[A-Z]+", ref).group(0)
    index = lambda col: sum((ord(ch) - 64) * 26 ** k for k, ch in enumerate(reversed(col))) - 1
    out_rows = []
    for row in re.findall(r"<row[^>]*>(.*?)</row>", z.read("xl/worksheets/sheet1.xml").decode(), re.S):
        cells = {}
        for ref, attrs, body in re.findall(r'<c r="([A-Z]+\d+)"([^>]*?)(?:/>|>(.*?)</c>)', row, re.S):
            v = re.search(r"<v>(.*?)</v>", body or "")
            v = v.group(1) if v else html.unescape(re.sub(r"<[^>]+>", "", body or ""))
            if 't="s"' in attrs and v != "":
                v = ss[int(v)]
            elif v != "" and 't="' not in attrs:
                v = float(v) if "." in v else int(v)
            cells[index(letters(ref))] = v
        out_rows.append([cells.get(i, "") for i in range(max(cells) + 1)] if cells else [])
    width = len(out_rows[0])
    return [r + [""] * (width - len(r)) for r in out_rows]


rows = read_rows(src)
header = [h for h in rows[0]]
drop = {"TAM", "SAM", "SOM"}
keep = [i for i, h in enumerate(header) if h not in drop]
cols = [header[i] for i in keep]
col = {h: n for n, h in enumerate(cols)}

data = [[r[i] for i in keep] for r in rows[1:16]]          # first 15 projects, same OED IDs as dataset2
by_name = {r[col["Platform Name"]]: r for r in data}

# A few edits so the compare view has something to show besides stages
edits = [
    ("Nam Du FPSO", "Hull award year", 2027),             # was 2026 → slipped 1 year
    ("Bonga Southwest - Aparo FPSO", "Installation year", 2033),   # was 2031 → slipped 2 years
    ("Fangtooth FPSO (FPSO-9)", "Hull award year", 2026),  # was 2027 → 1 year earlier
    ("Bay du Nord FPSO", "Facility Status", "Bidding"),    # was FEED
    ("Abadi FPSO", "EPCI Capex", 2700),                    # was 2550
]
changed = []
for name, field, value in edits:
    if name in by_name:
        before = by_name[name][col[field]]
        by_name[name][col[field]] = value
        changed.append(f"{name}: {field} {before} -> {value}")

def new(oed, name, ftype, status, product, loc, country, region, award, install, hull, owned, operator,
        contractor, green, depth, topside_t, config, capex, oil, gas):
    row = [""] * len(cols)
    vals = {"OED ID": oed, "Platform Name": name, "Facility Type": ftype, "Facility Status": status, "product": product,
            "Location": loc, "Country": country, "WMSuperRegion": region, "Hull award year": award,
            "Installation year": install, "New Build Hull": hull, "OWNED/LEASED": owned, "Operator Group": operator,
            "EPCI Floater contractor": contractor, "Platform Green/Brownfield": green, "Water Depth Group (M)": depth,
            "Topside Weight (Tonnes)": topside_t, "Topside Configuration": config, "EPCI Capex": capex,
            "Max Oil Production Rate (bbl/d)": oil, "Max Gas Production Rate (mmcfd)": gas}
    for k, v in vals.items():
        if k in col:
            row[col[k]] = v
    return row

# Five made-up projects (names say TEST so nobody mistakes them for real ones)
data += [
    new(990001, "TEST Atlantis FPSO", "FPSO", "Conceptual", "Oil/Gas", "24 10S 42 05W", "Brazil",
        "Latin America and the Caribbean", 2027, 2031, "New Build", "Leased", "Petrobras", "", "Greenfield",
        "1500-1999", 32000, "Acc/Prod/Sep/Inj/Sto/Fl", 2100, 180000, 250),
    new(990002, "TEST Kora FPSO", "FPSO", "FEED", "Oil/Gas", "5 12S 11 40E", "Angola", "Africa", 2026, 2029,
        "Conversion", "OpOwn", "TotalEnergies", "", "Greenfield", "1000-1499", 18000, "Acc/Prod/Sep/Inj/Sto/Fl",
        1250, 100000, 120),
    new(990003, "TEST Nordlys FPSO", "FPSO", "Conceptual", "Oil", "71 30N 21 10E", "Norway", "Europe", 2028, 2032,
        "New Build", "OpOwn", "Equinor", "", "Greenfield", "100-499", 22000, "Acc/Prod/Sep/Sto/Fl", 1600, 110000, 0),
    new(990004, "TEST Hammerhead South FPSO", "FPSO", "Bidding", "Oil/Gas", "7 45N 56 50W", "Guyana",
        "Latin America and the Caribbean", 2026, 2029, "New Build", "Leased", "ExxonMobil", "SBM Offshore",
        "Greenfield", "1500-1999", 35000, "Acc/Prod/Sep/Inj/Sto/Fl", 2400, 220000, 400),
    new(990005, "TEST Sabah Deep FPSO", "FPSO", "Conceptual", "Gas/Condensate", "6 20N 115 05E", "Malaysia", "Asia",
        2028, 2031, "Conversion", "OpOwn", "Shell", "", "Greenfield", "1000-1499", 9000, "Acc/Prod/Sep/Sto/Fl",
        520, 0, 450),
]

wb = Workbook()
ws = wb.active
ws.title = "Projects"
ws.append(cols)
for r in data:
    ws.append(r)
bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0B1B33")
for c in ws[1]:
    c.font, c.fill = bold, fill
for i, h in enumerate(cols, start=1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(40, len(str(h)) + 4))
ws.column_dimensions[ws.cell(1, col["Platform Name"] + 1).column_letter].width = 36
ws.freeze_panes = "C2"
wb.save(out)

print(f"{len(data)} projects, {len(cols)} columns, no TAM/SAM/SOM")
print("Edited vs dataset2:", *changed, sep="\n  ")
print("Not carried over from dataset2:", *[r[header.index('Platform Name')] for r in rows[16:]], sep="\n  ")
